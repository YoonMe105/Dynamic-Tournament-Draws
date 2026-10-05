"""
Reads a ranking list from an Excel (.xlsx / .xls) or CSV file.

Spreadsheets have no fixed layout, so the columns are found from the
header row by name:

    rank         "Rank", "Position", "Pos", "Ranking"
                 or separate typed columns, used for that ranking only:
                 "National Ranking", "AJSS Ranking" / "Asian Ranking",
                 "World Ranking" / "PSA Ranking"  (e.g. a tournament player list)
    name         "Name", "Player", "Player Name"  (or "First Name" + "Last Name")
    dob          "DOB", "D.O.B", "Date of Birth", "Birth"
    ASF number   "Member ID", "ASF", "AJSS Member ID", "Member No"
    Rankedin ID  "Rankedin", "Rankedin ID"
    category     "Category", "Cat", "Division", "Event"
    country      "Country", "Nationality"
    state        "State"

The category can also come from the sheet name ("BU15") or a title line
("BOYS UNDER 15") when there's no Category column. Every sheet is read.

    parse_spreadsheet("ranking.xlsx") -> {"rows": [...], "text": "..."}
"""

import csv
import re
from datetime import date, datetime
from pathlib import Path

CATEGORY_RE = re.compile(r"\b([BG])U\s*0?(\d{1,2})\b", re.IGNORECASE)
CATEGORY_WORDS_RE = re.compile(r"(BOYS|GIRLS)\s+UNDER\s+0?(\d{1,2})", re.IGNORECASE)

# header text (lower case, letters only) -> field. First match wins.
HEADER_KEYWORDS = [
    ("rankedin", ["rankedin"]),
    ("asf", ["asf", "memberid", "memberno", "membershipno", "ajssmemberid", "ajssid", "membershipid", "ajssmembershipid"]),
    ("dob", ["dob", "dateofbirth", "birthdate", "birth"]),
    ("first_name", ["firstname"]),
    ("last_name", ["lastname", "surname"]),
    ("name", ["playername", "name", "player"]),
    # Typed ranking columns come before the plain "rank" one
    ("rank_ajss", ["ajssranking", "ajssrank", "asianranking", "asianjuniorranking"]),
    ("rank_world", ["worldranking", "worldrank", "psaranking", "psarank"]),
    ("rank_national", ["nationalranking", "nationalrank", "sramranking", "sramrank"]),
    ("rank", ["rank", "ranking", "position", "pos"]),
    ("category", ["category", "cat", "division", "event", "agegroup"]),
    ("country", ["country", "nationality", "nation"]),
    ("state", ["state"]),
]

# Read at most this many rows looking for the header
HEADER_SEARCH_ROWS = 30


def _key(text) -> str:
    return re.sub(r"[^a-z]", "", str(text or "").lower())


def _field_for(header) -> str | None:
    key = _key(header)

    if not key:
        return None

    for field, words in HEADER_KEYWORDS:

        # "Player ID", "Member ID" ... are IDs, never the name/rank column
        if key.endswith("id") and field not in ("asf", "rankedin"):
            continue

        for word in words:
            # exact for short words ("cat", "pos"), contains for longer ones
            if key == word or (len(word) > 3 and word in key):
                return field

    return None


def category_from_text(text: str) -> str | None:
    m = CATEGORY_RE.search(text or "")
    if m:
        return f"{m.group(1).upper()}U{int(m.group(2)):02d}"

    m = CATEGORY_WORDS_RE.search(text or "")
    if m:
        return f"{'B' if m.group(1).upper() == 'BOYS' else 'G'}U{int(m.group(2)):02d}"

    return None


def to_iso_date(value) -> str | None:
    """Excel date, or text like 24/11/2011, 24-11-11, 2011-11-24 -> yyyy-mm-dd."""
    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    text = str(value or "").strip()

    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
    if m:
        year, month, day = map(int, m.groups())
        return _safe_date(year, month, day)

    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})$", text)
    if m:
        day, month, year = map(int, m.groups())
        if year < 100:
            year += 2000 if 2000 + year <= date.today().year else 1900
        return _safe_date(year, month, day)

    return None


def _safe_date(year, month, day) -> str | None:
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return None


def to_rank(value) -> int | None:
    """12, "12", "12.0", "T12", "=12" -> 12"""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value) if value > 0 else None

    m = re.search(r"\d+", str(value or ""))
    return int(m.group()) if m and int(m.group()) > 0 else None


def clean_id(value) -> str:
    """Member numbers saved as numbers come back as 2217272.0 - drop the .0"""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value or "").strip()


# ----------------------------------------------------------------------
# Reading the file
# ----------------------------------------------------------------------

def _sheets(path: Path):
    """Yields (sheet name, list of rows) for every sheet."""
    suffix = path.suffix.lower()

    if suffix == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            yield path.stem, list(csv.reader(f))

    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))

        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    else:
                        row.append(cell.value)
                rows.append(row)
            yield sheet.name, rows

    else:
        import openpyxl

        book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)

        for sheet in book.worksheets:
            yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]

        book.close()


RANK_FIELDS = ("rank", "rank_ajss", "rank_world", "rank_national")


class AmbiguousColumn(Exception):
    """The typed column name fits more than one heading."""

    def __init__(self, wanted: str, headings: list):
        super().__init__(wanted)
        self.wanted = wanted
        self.headings = headings


def _column_named(row: list, wanted: str) -> int | None:
    """Column whose header is `wanted` (ignoring case, spaces and punctuation).
    A header that merely contains it is only accepted when exactly one does -
    "Rank" next to both "National Ranking" and "AJSS Ranking" is ambiguous."""
    target = _key(wanted)
    keys = [_key(cell) for cell in row]

    if not target:
        return None

    for col, key in enumerate(keys):
        if key and key == target:
            return col

    partial = [col for col, key in enumerate(keys) if key and target in key]

    if len(partial) > 1:
        raise AmbiguousColumn(wanted, [str(row[col]).strip() for col in partial])

    return partial[0] if partial else None


def _find_header(rows: list, rank_column: str | None = None) -> tuple[int, dict] | None:
    """Index of the header row and {field: column} - needs a rank and a name column.

    With rank_column, that column (by its header text) is the only ranking
    column used, whatever it's called."""
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        columns = {}

        for col, cell in enumerate(row):
            field = _field_for(cell)
            if field and field not in columns:
                columns[field] = col

        if rank_column:
            chosen = _column_named(row, rank_column)

            if chosen is None:
                continue

            for field in RANK_FIELDS:
                columns.pop(field, None)

            columns["rank"] = chosen

        has_name = "name" in columns or "last_name" in columns or "first_name" in columns
        has_rank = any(field in columns for field in RANK_FIELDS)

        if has_rank and has_name:
            return index, columns

    return None


def _header_names(rows: list) -> list:
    """Header texts of the first row that has a name column - for error messages."""
    for row in rows[:HEADER_SEARCH_ROWS]:
        fields = {_field_for(cell) for cell in row}

        if fields & {"name", "first_name", "last_name"}:
            return [str(cell).strip() for cell in row if cell not in (None, "") and str(cell).strip()]

    return []


# typed column field -> key in each row's "ranks"
TYPED_RANK_FIELDS = {"rank": "plain", "rank_ajss": "ajss", "rank_national": "national", "rank_world": "world"}


def parse_spreadsheet(path: str, rank_column: str | None = None) -> dict:
    """Returns {"rows": [...], "text": all title text (for detecting the ranking type),
    "rankColumns": which kinds of ranking column the file has ("plain", "ajss", "national", "world")}.

    Each row has "ranks": {"plain": 12, "ajss": None, ...} - check_rankings picks the
    one that matches the ranking being checked.

    rank_column: header of the one column to read the ranking from (chosen by
    the admin); every other ranking column is ignored."""
    path = Path(path)
    rows_out = []
    title_text = []
    rank_columns = set()
    headers_seen = []

    for sheet_name, rows in _sheets(path):
        header = _find_header(rows, rank_column)

        if header is None:
            headers_seen.extend(h for h in _header_names(rows) if h not in headers_seen)
            continue

        header_index, columns = header

        # Text above the header: titles like "ASIAN JUNIOR RANKING ..." / "BOYS UNDER 15"
        above = " ".join(str(c) for row in rows[:header_index] for c in row if c not in (None, ""))
        title_text.append(f"{sheet_name} {above}")

        sheet_category = category_from_text(sheet_name) or category_from_text(above)

        for field, kind in TYPED_RANK_FIELDS.items():
            if field in columns:
                rank_columns.add(kind)

        def cell(row, field):
            col = columns.get(field)
            return row[col] if col is not None and col < len(row) else None

        for row in rows[header_index + 1:]:
            ranks = {kind: to_rank(cell(row, field)) for field, kind in TYPED_RANK_FIELDS.items() if field in columns}

            if "name" in columns:
                name = str(cell(row, "name") or "").strip()
            else:
                name = f"{cell(row, 'first_name') or ''} {cell(row, 'last_name') or ''}".strip()

            if not name or all(rank is None for rank in ranks.values()):
                continue   # blank line, sub-heading, total row, or a player with no ranking at all

            # Category column wins over the sheet name / title
            row_category = category_from_text(str(cell(row, "category") or "")) if "category" in columns else None

            rows_out.append({
                "ranks": ranks,
                "name": re.sub(r"\s+", " ", name),
                "dob": to_iso_date(cell(row, "dob")),
                "asfMemberNo": clean_id(cell(row, "asf")),
                "rankedinId": clean_id(cell(row, "rankedin")),
                "category": row_category or sheet_category,
                "country": str(cell(row, "country") or "").strip(),
                "state": str(cell(row, "state") or "").strip(),
            })

    return {"rows": rows_out, "text": " ".join(title_text), "rankColumns": rank_columns, "headers": headers_seen}
